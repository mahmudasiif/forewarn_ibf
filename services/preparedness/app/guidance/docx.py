"""Word (.docx) export — ported from src/lib/docxTemplate.ts to python-docx.

Renders the GuidelineDocument (current/forecast state, impact summary table,
major-impact sector lists, advisories by level) into a .docx and returns the
bytes. Bengali runs are tagged with a Bengali-capable complex-script font so
Word renders them correctly.
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

from app.guidance.schema import LEVEL_KEYS, impact_table_headers

SECTOR_LABELS = [
    ("Population and House", "populationAndHouse"),
    ("Infrastructure/Polder", "infrastructurePolder"),
    ("Agriculture", "agriculture"),
    ("Livestock + Fisheries", "livestockFisheries"),
    ("Livelihood", "livelihood"),
    ("Health/Disease", "healthDisease"),
    ("Utility service", "utilityService"),
]

TABLE_WIDTHS = [1350, 650, 650, 800, 1150, 650, 650, 650, 750, 800, 650, 610]
_BENGALI = re.compile(r"[ঀ-৿]")
_LATIN_FONT = "Arial"
_BENGALI_FONT = "Bangla Sangam MN"


def _set_run_font(run, name: str) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _add_runs(paragraph, text: str, *, bold: bool = False, color: str | None = None) -> None:
    for seg in re.split(r"([ঀ-৿]+)", text or ""):
        if not seg:
            continue
        run = paragraph.add_run(seg)
        run.bold = bold
        _set_run_font(run, _BENGALI_FONT if _BENGALI.search(seg) else _LATIN_FONT)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def _para(doc, text, *, style=None, bold=False, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    _add_runs(p, text, bold=bold)
    return p


def _heading(doc, text, level):
    # Use plain paragraphs with manual sizing so Bengali fonts apply cleanly.
    sizes = {0: 20, 1: 15, 2: 13}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8 if level else 0)
    p.paragraph_format.space_after = Pt(6)
    for seg in re.split(r"([ঀ-৿]+)", text or ""):
        if not seg:
            continue
        run = p.add_run(seg)
        run.bold = True
        run.font.size = Pt(sizes.get(level, 13))
        _set_run_font(run, _BENGALI_FONT if _BENGALI.search(seg) else _LATIN_FONT)
    return p


def _field_line(doc, label, value):
    p = doc.add_paragraph(style="List Bullet")
    _add_runs(p, f"{label}: ", bold=True)
    _add_runs(p, value or "")


def _bullets(doc, items):
    if not items:
        _para(doc, "তথ্য নেই")
        return
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        _add_runs(p, item)


def _shade_cell(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def build_guideline_docx(doc: dict, track_map_data: bytes | None = None) -> bytes:
    d = Document()
    d.styles["Normal"].font.name = _LATIN_FONT
    d.styles["Normal"].font.size = Pt(11)

    _heading(d, f"Cyclone IBF Guideline — {doc.get('cycloneName', '')}", 0)

    warnings = doc.get("dataQualityWarnings") or []
    if warnings:
        _heading(d, "Data quality warning", 1)
        _bullets(d, warnings)

    cs = doc.get("currentState", {})
    _heading(d, "Current State", 1)
    _field_line(d, "Name of the TC", cs.get("name", ""))
    _field_line(d, "Category (IMD)", cs.get("category", ""))
    _field_line(d, "Wind Speed (Km/hr)", cs.get("windSpeedKmh", ""))
    _field_line(d, "Direction of the system", cs.get("direction", ""))
    _field_line(d, "Location", cs.get("location", ""))

    fs = doc.get("forecastState", {})
    _heading(d, "Forecast State", 1)
    _field_line(d, "Landfall location (upazila and union)", fs.get("landfallLocation", ""))
    _field_line(d, "Wind speed (max during/after landfall)", fs.get("windSpeed", ""))
    _field_line(d, "Storm surge (max during/after)", fs.get("stormSurge", ""))
    _field_line(d, "Time", fs.get("time", ""))
    _field_line(d, "Category (max during/after)", fs.get("category", ""))
    _field_line(d, "Precipitation (max during/after)", fs.get("precipitation", ""))

    if track_map_data:
        _heading(d, "Model-derived Track Map", 2)
        try:
            pic_p = d.add_paragraph()
            pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic_p.add_run().add_picture(io.BytesIO(track_map_data), width=Twips(9360))
        except Exception:  # noqa: BLE001 - a bad image must not fail the whole export
            pass

    _heading(d, "Cyclone Impact", 1)
    _heading(d, "Summary table", 2)

    headers = impact_table_headers(doc.get("rainfallWindowHours"))
    rows = doc.get("impactSummaryTable") or []
    table = d.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for i, cell in enumerate(table.rows[0].cells):
        cell.width = Twips(TABLE_WIDTHS[i] if i < len(TABLE_WIDTHS) else 700)
        _shade_cell(cell, "0369A1")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_runs(cell.paragraphs[0], headers[i], bold=True, color="FFFFFF")

    _keys = [
        "area", "windSpeedKmh", "surgeM", "precipitationMm", "totalPopulation",
        "affectedPopulation", "affectedMale", "affectedFemale", "affectedAge0To4",
        "affectedAge5To19", "affectedAge60Plus", "affectedHouseStructure",
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, key in enumerate(_keys):
            cells[i].width = Twips(TABLE_WIDTHS[i] if i < len(TABLE_WIDTHS) else 700)
            _add_runs(cells[i].paragraphs[0], str(row.get(key, "")))

    _heading(d, "Major Impact", 2)
    major = doc.get("majorImpact", {})
    for label, key in SECTOR_LABELS:
        _field_line(d, label, "")
        _bullets(d, major.get(key) or [])

    _heading(d, "Advisories", 1)
    advisories = doc.get("advisories", {})
    level_labels = doc.get("levelLabels", {})
    for level in LEVEL_KEYS:
        _para(d, level_labels.get(level, level), bold=True)
        sectors = advisories.get(level) or {}
        for label, key in SECTOR_LABELS:
            _field_line(d, label, "")
            _bullets(d, sectors.get(key) or [])

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
