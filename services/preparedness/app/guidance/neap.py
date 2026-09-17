"""NEAP sector-wise action context — ported from src/lib/neap/{actions,xlsxReader}.ts.

Reads the four NEAP action spreadsheets (openpyxl instead of the TS hand-rolled
xlsx parser), routes each sector/subsector row to impact categories, and renders
the readiness/early-action block for the guideline prompt, keyed on lead time to
landfall (48h threshold).
"""
from __future__ import annotations

import os
import re

from openpyxl import load_workbook

from app.guidance.schema import CATEGORY_LABELS
from app.settings import settings

ACTIONS_DIR = os.path.join(settings.ASSETS_DIR, "NEAP sector-wise actions spreadsheet-selected")

EARLY_ACTION_THRESHOLD_HOURS = 48

_COLUMN_KEYS = (
    "sector",
    "subsector",
    "communityReadiness",
    "institutionalReadiness",
    "communityEarly",
    "institutionalEarly",
)

_HEADER_MATCHERS: list[tuple[str, "re.Pattern[str]"]] = [
    ("sector", re.compile(r"^sector$")),
    ("subsector", re.compile(r"^sub[\s-]?sector$")),
    ("communityReadiness", re.compile(r"(?=.*readiness)(?=.*(communit|c-ra))")),
    ("institutionalReadiness", re.compile(r"(?=.*readiness)(?=.*(institution|i-ra))")),
    ("communityEarly", re.compile(r"(?=.*early action)(?=.*(communit|c-ea))")),
    ("institutionalEarly", re.compile(r"(?=.*early action)(?=.*(institution|i-ea))")),
]

_NOT_APPLICABLE = re.compile(r"^n\s*/?\s*a\.?$", re.IGNORECASE)
_ENUMERATOR = re.compile(r"^(?:\d{1,2}[.)]|[a-zA-Z][.)]|[-–—•*])\s+")

_CATEGORY_KEYWORDS: list[tuple[str, "re.Pattern[str]"]] = [
    ("livestockFisheries", re.compile(r"livestock|cattle|poultry|killa|fisher|shrimp|fish|pond|aquacult|veterinar")),
    ("agriculture", re.compile(r"agricultur|crop|seedbed|seed bed|salt farm|forest|tree|vegetation|horticult")),
    ("healthDisease", re.compile(r"health|wash|sanitation|hygiene|disease|medical|medicine|nutrition")),
    ("utilityService", re.compile(r"utility|telecom|electric|network|power|transport|water")),
    ("infrastructurePolder", re.compile(r"infrastructure|polder|embankment|sluice|shelter management|school|education|building|structur|road")),
    ("livelihood", re.compile(r"livelihood|wage|expenditure|cash|income|food security|market|employment")),
    ("populationAndHouse", re.compile(r"life|asset|protection|gender|gbv|child|household|house|shelter|population|people|document|warning|communication|disab|older|elder|inclusion|evacuation|safety")),
]
_MAX_CATEGORIES_PER_ROW = 2


def _norm_header(cell: str) -> str:
    return re.sub(r"\s+", " ", cell or "").strip().lower()


def _resolve_columns(header_row: list[str]) -> dict[str, int] | None:
    col: dict[str, int] = {}
    matched = 0
    for index, cell in enumerate(header_row):
        h = _norm_header(cell)
        if not h:
            continue
        for key, test in _HEADER_MATCHERS:
            if key not in col and test.search(h):
                col[key] = index
                matched += 1
                break
    if "sector" not in col or "subsector" not in col or matched < 3:
        return None
    return col


def _split_actions(cell: str | None) -> list[str]:
    if not cell:
        return []
    out = []
    for line in re.split(r"\r?\n", cell):
        line = re.sub(r"\s+", " ", line).strip()
        if not line or _NOT_APPLICABLE.match(line):
            continue
        line = _ENUMERATOR.sub("", line).strip()
        if line:
            out.append(line)
    return out


def _match_categories(text: str) -> list[str]:
    haystack = (text or "").lower()
    hits: list[str] = []
    for category, pattern in _CATEGORY_KEYWORDS:
        if pattern.search(haystack):
            hits.append(category)
        if len(hits) == _MAX_CATEGORIES_PER_ROW:
            break
    return hits


def _route_categories(sector: str, subsector: str) -> list[str]:
    from_sub = _match_categories(subsector)
    return from_sub if from_sub else _match_categories(sector)


_cache: dict | None = None


def _spreadsheet_paths() -> list[str]:
    try:
        names = os.listdir(ACTIONS_DIR)
    except OSError as exc:  # noqa: BLE001
        print(f"[NEAP] cannot read {ACTIONS_DIR}: {exc}", flush=True)
        return []
    return [
        os.path.join(ACTIONS_DIR, n)
        for n in sorted(names)
        if n.lower().endswith(".xlsx") and not n.startswith("~$")
    ]


def _signature(files: list[str]) -> str:
    parts = []
    for f in files:
        try:
            s = os.stat(f)
            parts.append(f"{os.path.basename(f)}:{s.st_mtime}:{s.st_size}")
        except OSError:
            parts.append(f"{os.path.basename(f)}:missing")
    return "|".join(parts)


def _read_sheets(path: str) -> list[tuple[str, list[list[str]]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            rows.append(cells)
        sheets.append((ws.title, rows))
    wb.close()
    return sheets


def load_neap_action_rows() -> list[dict]:
    global _cache
    files = _spreadsheet_paths()
    sig = _signature(files)
    if _cache and _cache["signature"] == sig:
        return _cache["rows"]

    rows: list[dict] = []
    for file in files:
        name = os.path.basename(file)
        try:
            for sheet_name, sheet_rows in _read_sheets(file):
                header_index = next(
                    (i for i, r in enumerate(sheet_rows) if _resolve_columns(r) is not None), -1
                )
                if header_index < 0:
                    continue
                columns = _resolve_columns(sheet_rows[header_index])
                for r in sheet_rows[header_index + 1 :]:
                    def cell(key: str) -> str:
                        idx = columns[key]
                        return r[idx] if idx < len(r) else ""

                    sector = re.sub(r"\s+", " ", cell("sector")).strip()
                    subsector = re.sub(r"\s+", " ", cell("subsector")).strip()
                    if not sector and not subsector:
                        continue
                    parsed = {
                        "file": name,
                        "sheet": sheet_name,
                        "sector": sector,
                        "subsector": "" if _NOT_APPLICABLE.match(subsector) else subsector,
                        "communityReadiness": _split_actions(cell("communityReadiness")),
                        "institutionalReadiness": _split_actions(cell("institutionalReadiness")),
                        "communityEarly": _split_actions(cell("communityEarly")),
                        "institutionalEarly": _split_actions(cell("institutionalEarly")),
                        "categories": _route_categories(sector, subsector),
                    }
                    if any(
                        parsed[k]
                        for k in ("communityReadiness", "institutionalReadiness", "communityEarly", "institutionalEarly")
                    ):
                        rows.append(parsed)
        except Exception as exc:  # noqa: BLE001
            print(f"[NEAP] failed to parse {name}: {exc}", flush=True)

    _cache = {"signature": sig, "rows": rows}
    print(f"[NEAP] loaded {len(rows)} sector rows from {len(files)} spreadsheet(s)", flush=True)
    return rows


def resolve_phase(lead_time_hours) -> str:
    if not isinstance(lead_time_hours, (int, float)) or lead_time_hours != lead_time_hours:
        return "both"
    return "readiness" if lead_time_hours > EARLY_ACTION_THRESHOLD_HOURS else "early"


_READINESS = {
    "community": "communityReadiness",
    "institutional": "institutionalReadiness",
    "label": {
        "en": "READINESS ACTION (RA) — more than 48 hours to landfall",
        "bn": "প্রস্তুতিমূলক কার্যক্রম (RA) — ল্যান্ডফলের ৪৮ ঘণ্টার বেশি বাকি",
    },
}
_EARLY = {
    "community": "communityEarly",
    "institutional": "institutionalEarly",
    "label": {
        "en": "EARLY ACTION (EA) — 48 hours or less to landfall",
        "bn": "আগাম কার্যক্রম (EA) — ল্যান্ডফলের ৪৮ ঘণ্টা বা তার কম বাকি",
    },
}


def _phase_columns(phase: str) -> list[dict]:
    if phase == "readiness":
        return [_READINESS]
    if phase == "early":
        return [_EARLY]
    return [_READINESS, _EARLY]


def _row_title(row: dict) -> str:
    return f"{row['sector']} › {row['subsector']}" if row["subsector"] else row["sector"]


def _build_inventory(rows: list[dict], lang: str) -> str:
    by_file: dict[str, list[str]] = {}
    for row in rows:
        if row["categories"]:
            categories = " + ".join(CATEGORY_LABELS[lang][c] for c in row["categories"])
        else:
            categories = "Advisories only" if lang == "en" else "শুধু পরামর্শ"
        by_file.setdefault(row["file"], []).append(f"{_row_title(row)} → {categories}")
    return "\n".join(f"- **{file}**: {'; '.join(entries)}" for file, entries in by_file.items())


def _render_row(row: dict, columns: dict, lang: str) -> str:
    community = row[columns["community"]]
    institutional = row[columns["institutional"]]
    if not community and not institutional:
        return ""
    if row["categories"]:
        categories = " + ".join(CATEGORY_LABELS[lang][c] for c in row["categories"])
    else:
        categories = "Advisories" if lang == "en" else "পরামর্শ"
    target_line = f"Template section: {categories}" if lang == "en" else f"টেমপ্লেট বিভাগ: {categories}"
    community_label = "Community level" if lang == "en" else "কমিউনিটি পর্যায়"
    institutional_label = "Institutional level" if lang == "en" else "প্রাতিষ্ঠানিক পর্যায়"
    lines = [f"#### {_row_title(row)}", target_line]
    if community:
        lines.append(f"{community_label}:")
        lines.extend(f"- {item}" for item in community)
    if institutional:
        lines.append(f"{institutional_label}:")
        lines.extend(f"- {item}" for item in institutional)
    return "\n".join(lines)


def build_neap_action_context(lead_time_hours, lang: str) -> dict:
    rows = load_neap_action_rows()
    phase = resolve_phase(lead_time_hours)
    column_sets = _phase_columns(phase)

    item_count = 0
    sections: list[str] = []
    for columns in column_sets:
        rendered = []
        for row in rows:
            item_count += len(row[columns["community"]]) + len(row[columns["institutional"]])
            r = _render_row(row, columns, lang)
            if r:
                rendered.append(r)
        if rendered:
            sections.append(f"### {columns['label'][lang]}\n\n" + "\n\n".join(rendered))

    if isinstance(lead_time_hours, (int, float)) and lead_time_hours == lead_time_hours:
        lead_line = (
            f"Time to landfall: {lead_time_hours:.1f} hours."
            if lang == "en"
            else f"ল্যান্ডফল পর্যন্ত সময়: {lead_time_hours:.1f} ঘণ্টা।"
        )
    else:
        lead_line = (
            "Time to landfall could not be determined from the model track, so both phases are given below."
            if lang == "en"
            else "মডেল ট্র্যাক থেকে ল্যান্ডফল পর্যন্ত সময় নির্ধারণ করা যায়নি, তাই নিচে দুটি পর্যায়ই দেওয়া হলো।"
        )

    inventory = _build_inventory(rows, lang)
    block = f"{lead_line}\n\n{inventory}\n\n" + "\n\n".join(sections) if sections else ""
    return {
        "phase": phase,
        "leadTimeHours": lead_time_hours if isinstance(lead_time_hours, (int, float)) else None,
        "rows": rows,
        "inventory": inventory,
        "block": block,
        "itemCount": item_count,
    }
